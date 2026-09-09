"""从 docs/原子工艺清单v4.xlsx 重建原子工艺表及其关联表（atom/atom_alias/atom_category/dim_group）。

与 import_master_data.py 的区别：品类、工艺域、工艺阶段、工艺类别四张主数据表保持不变
（仅校验 v4 引用值必须已存在）；「合并映射表」中废弃编码的名称与别名作为别名并入合并目标原子，
保证历史报价叫法（如 去毛刺/清披锋）仍能匹配到合并后的原子。幂等：先清四张关联表再导入。
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

BACKEND_DIR = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND_DIR.parent
DEFAULT_XLSX = REPO_ROOT / "docs" / "原子工艺清单v4.xlsx"

sys.path.insert(0, str(BACKEND_DIR))
sys.path.insert(0, str(BACKEND_DIR / "scripts"))

from app.db import get_connection, init_db  # noqa: E402
from import_master_data import (  # noqa: E402
    domain_code_of,
    load_sheet,
    split_aliases,
    split_categories,
)

CLEAR_TABLES = ["atom_category", "atom_alias", "atom", "dim_group"]
MERGE_SHEET = "合并映射表"


def load_merge_rows(path: Path) -> list[tuple[str, str, str, str]]:
    """合并映射表 → (废弃名称, 废弃别名原文, 合并至编码, 合并至新编码)。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if MERGE_SHEET not in wb.sheetnames:
        return []
    rows = [tuple(r[:6]) for r in wb[MERGE_SHEET].iter_rows(values_only=True) if r and r[0]]
    return [
        (str(r[1]).strip(), str(r[2]).strip() if r[2] else "", str(r[3]).strip(), str(r[5]).strip())
        for r in rows[1:]
    ]


def main() -> None:
    xlsx_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        raise SystemExit(f"找不到文件：{xlsx_path}")

    rows = load_sheet(xlsx_path)
    merge_rows = load_merge_rows(xlsx_path)
    print(f"读取 {xlsx_path.name}：{len(rows)} 条原子，{len(merge_rows)} 条合并映射")

    init_db()
    conn = get_connection()
    try:
        db_domains = {n: c for c, n in conn.execute("SELECT code, name FROM process_domain")}
        db_stages = {n for (n,) in conn.execute("SELECT name FROM process_stage")}
        db_classes = {n for (n,) in conn.execute("SELECT name FROM process_class")}
        db_categories = {n: c for c, n in conn.execute("SELECT code, name FROM category")}

        atoms: list[tuple] = []
        aliases: list[tuple[str, str]] = []
        atom_categories: list[tuple[str, str]] = []
        seen_alias: set[tuple[str, str]] = set()

        def add_alias(atom_code: str, text: str) -> None:
            key = (atom_code, re.sub(r"\s+", "", text))
            if key in seen_alias:
                return
            seen_alias.add(key)
            aliases.append((atom_code, text))

        missing: dict[str, set[str]] = {"工艺域": set(), "工艺阶段": set(), "工艺类别": set(), "品类": set()}
        for code, name, alias_raw, domain, stage, klass, cats_raw, remark in rows:
            code = str(code).strip()
            name = str(name).strip()
            domain = str(domain).strip()
            stage = str(stage).strip()
            klass = str(klass).strip()
            if domain not in db_domains:
                missing["工艺域"].add(domain)
            if stage not in db_stages:
                missing["工艺阶段"].add(stage)
            if klass not in db_classes:
                missing["工艺类别"].add(klass)
            is_fallback = 1 if code == "AT-QT-001" else 0
            atoms.append((code, name, db_domains.get(domain, ""), stage, klass, remark, is_fallback))

            name_key = re.sub(r"\s+", "", name)
            for alias in split_aliases(alias_raw):
                if re.sub(r"\s+", "", alias) == name_key:
                    continue
                add_alias(code, alias)

            for cat_name in split_categories(cats_raw):
                cat_code = db_categories.get(cat_name)
                if cat_code is None:
                    missing["品类"].add(cat_name)
                else:
                    atom_categories.append((code, cat_code))

        if any(missing.values()):
            raise SystemExit(f"v4 引用了主数据表中不存在的值：{ {k: sorted(v) for k, v in missing.items() if v} }")

        # 废弃编码的名称/别名 → 合并目标原子的别名（历史报价叫法可继续匹配）
        atom_name_by_code = {a[0]: a[1] for a in atoms}
        alias_count_from_merge = 0
        for old_name, old_alias_raw, _old_code, new_code in merge_rows:
            if new_code not in atom_name_by_code:
                raise SystemExit(f"合并映射目标 {new_code} 不在综合原子清单中")
            target_name_key = re.sub(r"\s+", "", atom_name_by_code[new_code])
            for text in [old_name, *split_aliases(old_alias_raw)]:
                if re.sub(r"\s+", "", text) == target_name_key:
                    continue
                before = len(aliases)
                add_alias(new_code, text)
                alias_count_from_merge += len(aliases) - before

        with conn:
            for table in CLEAR_TABLES:
                conn.execute(f"DELETE FROM {table}")

            conn.executemany(
                """INSERT INTO atom
                   (code, name, domain_code, stage_name, class_name, remark, is_fallback)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                atoms,
            )
            conn.executemany(
                "INSERT INTO atom_alias (atom_code, alias_text, source) VALUES (?, ?, 'initial')",
                aliases,
            )
            conn.executemany(
                "INSERT INTO atom_category (atom_code, category_code) VALUES (?, ?)",
                atom_categories,
            )

            # 内置三维抽屉：域/阶段/类别各一个，member_atoms 为该值下的原子编码
            domain_names = {a[2]: dn for dn, dc in db_domains.items() for a in atoms if a[2] == dc}
            groups: list[tuple[str, str, str, str]] = []
            for dc in sorted({a[2] for a in atoms}):
                members = [a[0] for a in atoms if a[2] == dc]
                groups.append((f"builtin:domain:{dc}", domain_names.get(dc, dc), "process_domain",
                               json.dumps(members, ensure_ascii=False)))
            for stage in sorted({a[3] for a in atoms}):
                members = [a[0] for a in atoms if a[3] == stage]
                groups.append((f"builtin:stage:{stage}", stage, "process_stage",
                               json.dumps(members, ensure_ascii=False)))
            for klass in sorted({a[4] for a in atoms}):
                members = [a[0] for a in atoms if a[4] == klass]
                groups.append((f"builtin:class:{klass}", klass, "process_class",
                               json.dumps(members, ensure_ascii=False)))
            conn.executemany(
                """INSERT INTO dim_group (group_code, group_name, scope, member_atoms, is_builtin)
                   VALUES (?, ?, ?, ?, 1)""",
                groups,
            )

        for table in ["atom", "atom_alias", "atom_category", "dim_group"]:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            print(f"{table:16s} {count}")
        print(f"其中合并映射贡献别名：{alias_count_from_merge}")
        fallback = conn.execute("SELECT COUNT(*) FROM atom WHERE is_fallback=1").fetchone()[0]
        print(f"兜底原子 is_fallback=1：{fallback}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
