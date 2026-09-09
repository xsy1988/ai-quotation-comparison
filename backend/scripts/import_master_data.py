"""从 docs/原子工艺清单v2.xlsx 导入主数据（幂等：先清主数据表再导入，业务表与 supplier 不动）。"""

import json
import re
import sys
from pathlib import Path

import openpyxl

BACKEND_DIR = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND_DIR.parent
DEFAULT_XLSX = REPO_ROOT / "docs" / "原子工艺清单v2.xlsx"

sys.path.insert(0, str(BACKEND_DIR))

from app.db import get_connection, init_db  # noqa: E402

SHEET_NAME = "综合原子清单"

CATEGORIES: list[tuple[str, str]] = [
    ("CAT-WJWK", "五金外壳"),
    ("CAT-CMF", "CMF"),
    ("CAT-WJNZ", "五金内置"),
    ("CAT-SJ", "塑胶"),
    ("CAT-PCBA", "PCBA"),
    ("CAT-GJ", "硅胶"),
    ("CAT-BC", "包材"),
]

ALIAS_SPLIT_RE = re.compile(r"[/、]")
TRAILING_ETC_RE = re.compile(r"\s*等$")
CODE_RE = re.compile(r"^AT-([A-Z]+)-\d+$")


def split_aliases(raw: str | None) -> list[str]:
    if not raw:
        return []
    seen: set[str] = set()
    out: list[str] = []
    for part in ALIAS_SPLIT_RE.split(str(raw)):
        text = TRAILING_ETC_RE.sub("", part).strip()
        if not text:
            continue
        key = re.sub(r"\s+", "", text)
        if key in seen:
            continue
        seen.add(key)
        out.append(text)
    return out


def split_categories(raw: str | None) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[、，,]", str(raw))
    return [p.strip() for p in parts if p.strip()]


def clear_master_tables(conn) -> None:
    for table in [
        "atom_category",
        "atom_alias",
        "atom",
        "dim_group",
        "process_domain",
        "process_stage",
        "process_class",
        "category",
    ]:
        conn.execute(f"DELETE FROM {table}")


def load_sheet(path: Path) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"xlsx 中找不到工作表 {SHEET_NAME}，实际：{wb.sheetnames}")
    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    expected = ("编码", "原子", "别名/变体", "工艺域", "工艺阶段位置", "工艺类别", "常用品类", "备注")
    if tuple(header[:8]) != expected:
        raise SystemExit(f"表头不符：{header}")
    return [tuple(r[:8]) for r in rows[1:] if r and r[0]]


def domain_code_of(atom_code: str) -> str | None:
    m = CODE_RE.match(atom_code.strip())
    return m.group(1) if m else None


def main() -> None:
    xlsx_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.exists():
        raise SystemExit(f"找不到文件：{xlsx_path}")

    rows = load_sheet(xlsx_path)
    print(f"读取 {xlsx_path.name}：{len(rows)} 条原子")

    category_name_to_code = {name: code for code, name in CATEGORIES}
    unknown_categories: set[str] = set()

    domains: dict[str, str] = {}  # name -> code
    stages: set[str] = set()
    classes: set[str] = set()
    atoms: list[tuple] = []
    aliases: list[tuple[str, str]] = []
    atom_categories: list[tuple[str, str]] = []

    for code, name, alias_raw, domain, stage, klass, cats_raw, remark in rows:
        code = str(code).strip()
        name = str(name).strip()
        domain = str(domain).strip()
        stage = str(stage).strip()
        klass = str(klass).strip()
        is_fallback = 1 if code == "AT-QT-001" else 0

        if domain not in domains:
            dc = domain_code_of(code)
            if not dc:
                raise SystemExit(f"无法从编码 {code} 推导工艺域编码")
            domains[domain] = dc

        stages.add(stage)
        classes.add(klass)
        atoms.append((code, name, domains[domain], stage, klass, remark, is_fallback))

        name_key = re.sub(r"\s+", "", name)
        for alias in split_aliases(alias_raw):
            if re.sub(r"\s+", "", alias) == name_key:
                continue
            aliases.append((code, alias))

        for cat_name in split_categories(cats_raw):
            cat_code = category_name_to_code.get(cat_name)
            if cat_code is None:
                unknown_categories.add(cat_name)
            else:
                atom_categories.append((code, cat_code))

    if unknown_categories:
        raise SystemExit(f"常用品类列出现未知品类：{sorted(unknown_categories)}")

    init_db()
    conn = get_connection()
    try:
        with conn:
            clear_master_tables(conn)

            conn.executemany(
                "INSERT INTO category (code, name) VALUES (?, ?)", CATEGORIES
            )
            conn.executemany(
                "INSERT INTO process_domain (code, name) VALUES (?, ?)",
                [(dc, dn) for dn, dc in sorted(domains.items(), key=lambda x: x[1])],
            )
            conn.executemany(
                "INSERT INTO process_stage (name) VALUES (?)", [(s,) for s in sorted(stages)]
            )
            conn.executemany(
                "INSERT INTO process_class (name) VALUES (?)", [(c,) for c in sorted(classes)]
            )
            conn.executemany(
                """INSERT INTO atom
                   (code, name, domain_code, stage_name, class_name, remark, is_fallback)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                atoms,
            )
            conn.executemany(
                "INSERT INTO atom_alias (atom_code, alias_text, source) VALUES (?, ?, 'initial')",
                aliases,
            )
            conn.executemany(
                "INSERT INTO atom_category (atom_code, category_code) VALUES (?, ?)",
                atom_categories,
            )

            # 内置三维抽屉：每个域/阶段/类别值一个抽屉，member_atoms 为该值下的原子编码
            groups: list[tuple[str, str, str, str]] = []  # (code, name, scope, members_json)
            for domain_name, domain_code in domains.items():
                members = [a[0] for a in atoms if a[2] == domain_code]
                groups.append((f"builtin:domain:{domain_code}", domain_name, "process_domain", json.dumps(members, ensure_ascii=False)))
            for stage in sorted(stages):
                members = [a[0] for a in atoms if a[3] == stage]
                groups.append((f"builtin:stage:{stage}", stage, "process_stage", json.dumps(members, ensure_ascii=False)))
            for klass in sorted(classes):
                members = [a[0] for a in atoms if a[4] == klass]
                groups.append((f"builtin:class:{klass}", klass, "process_class", json.dumps(members, ensure_ascii=False)))
            conn.executemany(
                """INSERT INTO dim_group (group_code, group_name, scope, member_atoms, is_builtin)
                   VALUES (?, ?, ?, ?, 1)""",
                groups,
            )

        for table in [
            "category",
            "process_domain",
            "process_stage",
            "process_class",
            "atom",
            "atom_alias",
            "atom_category",
            "dim_group",
            "supplier",
        ]:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            print(f"{table:16s} {count}")
        fallback = conn.execute("SELECT COUNT(*) FROM atom WHERE is_fallback=1").fetchone()[0]
        print(f"兜底原子 is_fallback=1：{fallback}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
