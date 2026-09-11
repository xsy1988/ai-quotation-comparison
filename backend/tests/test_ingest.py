"""全格式接入测试：docx/pdf/image → IR、ingest_file 扩展名分派、查重门禁、OCR 路径。"""

import json
import sys
from pathlib import Path

import pytest

from app import ingest, ingest_formats
from app.ingest import excel_to_ir, ingest_file, sha256_of
from app.llm.client import LLMError, LLMUnavailable

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR / "scripts"))

FIXTURES = Path(__file__).resolve().parent / "fixtures"


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    """ingest 归档/IR 目录重定向到临时目录，不污染 backend/data。"""
    monkeypatch.setattr(ingest, "ARCHIVE_DIR", tmp_path / "archive")
    monkeypatch.setattr(ingest, "IR_DIR", tmp_path / "ir")
    return tmp_path


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------

def _make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("供应商报价单")
    doc.add_paragraph("供应商：深圳市样例五金有限公司")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "费用项目"
    table.cell(0, 1).text = "明细"
    table.cell(0, 2).text = "金额"
    table.cell(1, 0).text = "加工费"
    table.cell(1, 1).text = "CNC加工"
    table.cell(1, 2).text = "2.00"
    table.cell(2, 0).text = ""
    table.cell(2, 1).text = ""
    table.cell(2, 2).text = ""
    doc.save(path)
    return path


def test_docx_to_ir(sandbox):
    path = _make_docx(sandbox / "quote.docx")
    ir = ingest_formats.docx_to_ir(path, "hash")

    assert ir.file_type == "docx"
    assert [b.text for b in ir.blocks] == ["供应商报价单", "供应商：深圳市样例五金有限公司"]
    assert all(b.sheet == "doc" and b.col == 1 for b in ir.blocks)
    assert [t.sheet for t in ir.tables] == ["table_1", "table_1"]
    header, row = ir.tables
    assert header.row_number == 1
    assert row.row_number == 2  # 空行（第 3 行）被跳过
    assert [(c.col, c.value) for c in row.cells] == [(1, "加工费"), (2, "CNC加工"), (3, "2.00")]


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------

def _make_pdf(path: Path, with_text: bool) -> Path:
    reportlab = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    if with_text:
        c.drawString(72, 720, "Material Fee 4.50")
        c.drawString(72, 700, "Processing CNC 2.00")
    else:
        c.rect(72, 700, 100, 20)  # 只画图无文字 → 无文字层
    c.showPage()
    c.save()
    return path


def test_pdf_to_ir_with_text_layer(sandbox):
    path = _make_pdf(sandbox / "quote.pdf", with_text=True)
    ir = ingest_formats.pdf_to_ir(path, "hash")

    assert ir.file_type == "pdf"
    assert ir.sheets == ["page_1"]
    assert any("Material Fee" in b.text for b in ir.blocks)


def _ocr_chat_json_factory(calls: list, lines: str):
    def fake_chat_json(messages, *, model=None, client=None):
        calls.append(messages)
        return {"lines": lines}, {"model": model}

    return fake_chat_json


def test_pdf_to_ir_no_text_layer_uses_vision_ocr(sandbox, monkeypatch):
    """无文字层扫描页：pypdfium2 渲染 → vision OCR，结果归入 page_N sheet。"""
    calls = []
    monkeypatch.setattr(
        "app.llm.client.chat_json",
        _ocr_chat_json_factory(calls, "ocr|1|1:供应商报价单\nocr|2|1:加工费|2:2.00"),
    )
    path = _make_pdf(sandbox / "scan.pdf", with_text=False)

    ir = ingest_formats.pdf_to_ir(path, "hash")

    assert ir.file_type == "pdf"
    assert ir.sheets == ["page_1"]
    assert [(t.sheet, t.row_number, [(c.col, c.value) for c in t.cells]) for t in ir.tables] == [
        ("page_1", 1, [(1, "供应商报价单")]),
        ("page_1", 2, [(1, "加工费"), (2, "2.00")]),
    ]
    # 渲染后的 PNG 送进了 vision OCR
    content = calls[0][0]["content"]
    assert content[1]["type"] == "image_url"
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")


def test_pdf_to_ir_ocr_native_lines_clustered(sandbox, monkeypatch):
    """OCR 专用模型原生输出（rotate_rect/text 数组，无视提示词格式）按视觉行列聚类成 IR 行。

    rotate_rect 为 [x, y, h, w, angle]（angle 恒 90°，第三分量是文本高），行聚类按 y 区间重叠。
    """
    native = {
        "lines": [
            {"rotate_rect": [10, 10, 20, 120, 90], "text": "序号"},
            {"rotate_rect": [140, 12, 18, 80, 90], "text": "品名"},
            {"rotate_rect": [12, 40, 18, 10, 90], "text": "1"},
            {"rotate_rect": [142, 41, 18, 60, 90], "text": "F06外壳"},
        ]
    }

    def fake_chat_json(messages, *, model=None, client=None):
        return native, {"model": model}

    monkeypatch.setattr("app.llm.client.chat_json", fake_chat_json)
    path = _make_pdf(sandbox / "scan.pdf", with_text=False)

    ir = ingest_formats.pdf_to_ir(path, "hash")

    assert [(t.row_number, [(c.col, c.value) for c in t.cells]) for t in ir.tables] == [
        (1, [(1, "序号"), (2, "品名")]),
        (2, [(1, "1"), (2, "F06外壳")]),
    ]


def test_pdf_to_ir_ocr_native_stacked_header_aligns_with_body_column():
    """两行堆叠的表头单元格（脱墨/氧化）合并为完整列名并与表体同列号，版面 LLM 不再对错列。"""
    from app.ingest_formats import _cluster_native_lines

    items = [
        {"rotate_rect": [540, 200, 17, 21, 90], "text": "脱墨"},
        {"rotate_rect": [538, 218, 23, 23, 90], "text": "氧化"},
        {"rotate_rect": [280, 210, 17, 21, 90], "text": "材料"},
        {"rotate_rect": [280, 256, 13, 21, 90], "text": "2.20"},
        {"rotate_rect": [540, 256, 13, 21, 90], "text": "1.30"},
    ]
    text = _cluster_native_lines(items)
    body_row = next(l for l in text.splitlines() if ":2.20" in l)
    header_row = next(l for l in text.splitlines() if "脱墨" in l)
    body_col_130 = next(c.split(":")[0] for c in body_row.split("|") if c.endswith(":1.30"))
    header_cells = header_row.split("|")[2:]
    # 竖排碎片合并为一个完整单元格，列号与表体 1.30 对齐
    assert f"{body_col_130}:脱墨氧化" in header_cells


def test_pdf_to_ir_ocr_stacked_header_fragments_merge_full_column():
    """创锋回归：窄列竖排表头"镭雕/破氧/白"合并为完整列名，且不链式吞并紧贴的数据行。

    旧实现按运行 y_end 链式聚行：表头三段碎片与下方 0.40/0.30 数据行并入同一视觉行，
    再经全局列量化得到 17:破氧|17:白|17:0.40|17:镭雕 的伪列碎片，导致 LLM 张冠李戴。
    """
    from app.ingest_formats import _cluster_native_lines

    items = [
        {"rotate_rect": [10, 100, 16, 30, 90], "text": "项次"},
        {"rotate_rect": [60, 100, 16, 30, 90], "text": "品名"},
        {"rotate_rect": [500, 96, 16, 40, 90], "text": "镭雕"},
        {"rotate_rect": [502, 116, 16, 40, 90], "text": "破氧"},
        {"rotate_rect": [508, 136, 16, 20, 90], "text": "白"},
        {"rotate_rect": [560, 100, 16, 40, 90], "text": "全检"},
        # 数据行紧贴表头（与"白"底边仅 2px，小于旧链式 gap）：仍须独立成行
        {"rotate_rect": [10, 154, 16, 10, 90], "text": "1"},
        {"rotate_rect": [60, 154, 16, 60, 90], "text": "F06外壳"},
        {"rotate_rect": [500, 154, 16, 40, 90], "text": "0.40"},
        {"rotate_rect": [560, 154, 16, 40, 90], "text": "0.30"},
    ]
    lines = _cluster_native_lines(items).splitlines()

    assert len(lines) == 2  # 表头行 + 数据行，互不吞并
    header, body = lines
    header_cells = dict(p.split(":", 1) for p in header.split("|")[2:])
    body_cells = dict(p.split(":", 1) for p in body.split("|")[2:])
    # "镭雕破氧白"为完整列名，且与数据 0.40 同列；"全检"与 0.30 同列
    assert "镭雕破氧白" in header_cells.values()
    col_laser = next(c for c, v in header_cells.items() if v == "镭雕破氧白")
    col_qc = next(c for c, v in header_cells.items() if v == "全检")
    assert body_cells[col_laser] == "0.40"
    assert body_cells[col_qc] == "0.30"
    # 无伪列：同一行内列号唯一
    assert len(header_cells) == len(header.split("|")) - 2


def test_pdf_to_ir_multiline_header_cell_joined(sandbox):
    """pdfplumber 路径：换行表头单元格合并为完整列名，不带换行符拆散 IR 序列化。"""
    reportlab = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = sandbox / "multiline_header.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("STSong-Light", 10)
    x0, x1, x2 = 72, 250, 450
    ytop, ymid, ybot = 720, 680, 650
    for x in (x0, x1, x2):
        c.line(x, ytop, x, ybot)
    for y in (ytop, ymid, ybot):
        c.line(x0, y, x2, y)
    c.drawString(x0 + 4, ytop - 14, "工艺")
    c.drawString(x1 + 4, ytop - 14, "镭雕破氧")
    c.drawString(x1 + 4, ytop - 28, "白")
    c.drawString(x0 + 4, ymid - 14, "镭雕")
    c.drawString(x1 + 4, ymid - 14, "0.40")
    c.showPage()
    c.save()

    ir = ingest_formats.pdf_to_ir(path, "hash")

    values = [cell.value for t in ir.tables for cell in t.cells]
    assert "镭雕破氧白" in values  # 完整列名，不拆成伪列/碎片
    assert all("\n" not in str(v) for v in values)  # 无换行符破坏 IR 行序列化


def test_join_cell_lines_multiline_header():
    """单元格内换行合并：CJK 直接拼接，拉丁边界补空格。"""
    from app.ingest_formats import _join_cell_lines

    assert _join_cell_lines("镭雕\n破氧\n白") == "镭雕破氧白"
    assert _join_cell_lines("成品\n全检费") == "成品全检费"
    assert _join_cell_lines("合计价\n格") == "合计价格"
    assert _join_cell_lines("Unit\nPrice") == "Unit Price"


def test_pdf_to_ir_mixed_text_and_scan_pages(sandbox, monkeypatch):
    """混合文档：第 1 页文字层直接抽取，第 2 页扫描页走 OCR，各自归入 page_N。"""
    calls = []
    monkeypatch.setattr(
        "app.llm.client.chat_json",
        _ocr_chat_json_factory(calls, "ocr|1|1:盖章页"),
    )
    reportlab = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = sandbox / "mixed.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 720, "Material Fee 4.50")
    c.showPage()
    c.rect(72, 700, 100, 20)  # 第 2 页只画图无文字
    c.showPage()
    c.save()

    ir = ingest_formats.pdf_to_ir(path, "hash")

    assert ir.sheets == ["page_1", "page_2"]
    assert any("Material Fee" in b.text and b.sheet == "page_1" for b in ir.blocks)
    assert [(t.sheet, t.cells[0].value) for t in ir.tables] == [("page_2", "盖章页")]
    assert len(calls) == 1  # 只有第 2 页走了 OCR


# ---------------------------------------------------------------------------
# image（vision OCR，mock chat_json）
# ---------------------------------------------------------------------------

def _fake_png(path: Path) -> Path:
    path.write_bytes(b"\x89PNG\r\n\x1a\nfake-image-bytes")
    return path


def test_image_to_ir_prompt_and_multimodal(sandbox, monkeypatch):
    calls = {}

    def fake_chat_json(messages, *, model=None, client=None):
        calls["messages"] = messages
        calls["model"] = model
        lines = "ocr|1|1:供应商报价单\nocr|2|1:供应商|2:深圳样例公司\nocr|3|1:加工费|2:CNC加工|3:2.00"
        return {"lines": lines}, {"model": model}

    monkeypatch.setattr("app.llm.client.chat_json", fake_chat_json)
    path = _fake_png(sandbox / "scan.png")

    ir = ingest_formats.image_to_ir(path, "hash")

    assert calls["model"] == "qwen3.5-ocr"
    content = calls["messages"][0]["content"]
    assert content[0]["type"] == "text"
    assert "rotate_rect" in content[0]["text"] and '"lines"' in content[0]["text"]
    assert content[1]["type"] == "image_url"
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")

    assert ir.file_type == "image" and ir.sheets == ["ocr"]
    assert [(t.row_number, [(c.col, c.value) for c in t.cells]) for t in ir.tables] == [
        (1, [(1, "供应商报价单")]),
        (2, [(1, "供应商"), (2, "深圳样例公司")]),
        (3, [(1, "加工费"), (2, "CNC加工"), (3, "2.00")]),
    ]
    assert ir.blocks == []


def test_image_to_ir_unparseable_line_becomes_textblock(sandbox, monkeypatch):
    def fake_chat_json(messages, *, model=None, client=None):
        return {"lines": "ocr|1|1:正常行\n这是一行无法解析的内容"}, {}

    monkeypatch.setattr("app.llm.client.chat_json", fake_chat_json)
    ir = ingest_formats.image_to_ir(_fake_png(sandbox / "scan.jpg"), "hash")

    assert len(ir.tables) == 1 and ir.tables[0].row_number == 1
    assert len(ir.blocks) == 1 and ir.blocks[0].text == "这是一行无法解析的内容"


def test_image_to_ir_plain_text_compat(sandbox, monkeypatch):
    """网关忽略 response_format 时：dict 里没有 lines，取第一个字符串值按裸文本行处理。"""
    def fake_chat_json(messages, *, model=None, client=None):
        return {"text": "ocr|1|1:裸文本行"}, {}

    monkeypatch.setattr("app.llm.client.chat_json", fake_chat_json)
    ir = ingest_formats.image_to_ir(_fake_png(sandbox / "scan.png"), "hash")

    assert [(t.sheet, t.row_number, t.cells[0].value) for t in ir.tables] == [("ocr", 1, "裸文本行")]


def test_image_to_ir_llm_error_propagates(sandbox, monkeypatch):
    def failing(messages, *, model=None, client=None):
        raise LLMError("OCR 输出不是合法 JSON")

    monkeypatch.setattr("app.llm.client.chat_json", failing)
    with pytest.raises(LLMError):
        ingest_formats.image_to_ir(_fake_png(sandbox / "scan.png"), "hash")


def test_image_to_ir_llm_unavailable_propagates_no_fallback(sandbox, monkeypatch):
    """LLM 不可用：不降级，原样上抛（用户政策）。"""
    def failing(messages, *, model=None, client=None):
        raise LLMUnavailable("LLM 网关不可达或超时")

    monkeypatch.setattr("app.llm.client.chat_json", failing)
    with pytest.raises(LLMUnavailable):
        ingest_formats.image_to_ir(_fake_png(sandbox / "scan.png"), "hash")


# ---------------------------------------------------------------------------
# ingest_file 分派与查重门禁
# ---------------------------------------------------------------------------

def test_ingest_file_dispatches_by_suffix(sandbox, monkeypatch):
    calls = []
    fake_ir = ingest.excel_to_ir(FIXTURES / "sample_quote.xlsx", "fake")

    def fake_to_ir(path, file_hash, **kwargs):
        calls.append(path.suffix)
        return fake_ir

    for fn_name in ("docx_to_ir", "pdf_to_ir", "image_to_ir"):
        monkeypatch.setattr(ingest_formats, fn_name, fake_to_ir)

    src = FIXTURES / "sample_quote.xlsx"
    for i, suffix in enumerate((".docx", ".pdf", ".png")):
        path = sandbox / f"quote{i}{suffix}"  # 内容不同，避开 sha256 查重门禁
        path.write_bytes(src.read_bytes() + str(i).encode())
        result = ingest_file(path)
        assert result["status"] == "ingested"

    assert calls == [".docx", ".pdf", ".png"]


def test_ingest_file_unsupported_suffix(sandbox):
    path = sandbox / "quote.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持的文件格式"):
        ingest_file(path)


def test_ingest_file_dedup_gate_still_applies(sandbox):
    src = FIXTURES / "sample_quote.xlsx"
    first = ingest_file(src)
    assert first["status"] == "ingested"

    second = ingest_file(src)
    assert second["status"] == "reused"
    assert second["ir_path"] == first["ir_path"]

    forced = ingest_file(src, force=True)
    assert forced["status"] == "ingested"


def test_ingest_file_no_text_pdf_uses_ocr(sandbox, monkeypatch):
    """扫描版 PDF 经 ingest_file 全链路：渲染 + OCR 成功后正常归档登记。"""
    monkeypatch.setattr(
        "app.llm.client.chat_json",
        _ocr_chat_json_factory([], "ocr|1|1:供应商报价单"),
    )
    path = _make_pdf(sandbox / "scan.pdf", with_text=False)
    result = ingest_file(path)

    assert result["status"] == "ingested"
    assert result["sheets"] == ["page_1"]
    ir_data = json.loads(Path(result["ir_path"]).read_text(encoding="utf-8"))
    assert ir_data["file_type"] == "pdf"
    assert ir_data["tables"][0]["sheet"] == "page_1"


def test_excel_baseline_unchanged(sandbox):
    result = ingest_file(FIXTURES / "sample_quote.xlsx")
    assert result["status"] == "ingested"
    ir_data = json.loads(Path(result["ir_path"]).read_text(encoding="utf-8"))
    assert ir_data["file_type"] == "xlsx"
    assert ir_data["source_file"] == "sample_quote.xlsx"
    assert sha256_of(FIXTURES / "sample_quote.xlsx") == result["sha256"]
