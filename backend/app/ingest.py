"""文件接入（流水线段⓪①）：sha256 查重门禁 + 全格式 → IR + 原件归档 + IR 快照。

格式分派：.xlsx/.xlsm → excel_to_ir；.docx/.pdf → app.ingest_formats（结构化抽取，
PDF 扫描页自动渲染为图片走 vision OCR）；.png/.jpg/.jpeg → app.ingest_formats.image_to_ir（vision OCR，失败上抛不降级）。
查重门禁与归档/登记逻辑全格式共用。
"""

import hashlib
import json
import shutil
from pathlib import Path

import openpyxl

from .db import get_connection, init_db
from .formula_eval import solve_missing_formulas
from .ir import CellValue, IR, TableRow
from .normalize import display_number

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
ARCHIVE_DIR = DATA_DIR / "archive"
IR_DIR = DATA_DIR / "ir"

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".docx", ".pdf", ".png", ".jpg", ".jpeg"}


class ParseError(Exception):
    """文件解析失败（如无文字层 PDF、格式损坏），中文错误信息反馈给用户。"""


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_source_file(file_hash: str) -> dict | None:
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT * FROM source_file WHERE sha256 = ?", (file_hash,)
        ).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def formula_cell_to_coord(key: tuple[str, int, int]) -> str:
    """(表, 行, 列) → `表!R{行}C{列}`（与版面理解的位置写法一致，如 `报价单!R9C7`）。"""
    sheet, row, col = key
    return f"{sheet}!R{row}C{col}"


def excel_to_ir(path: Path, file_hash: str) -> IR:
    value_book = openpyxl.load_workbook(path, data_only=True)
    formula_book = openpyxl.load_workbook(path, data_only=False)
    solved, unresolved = solve_missing_formulas(value_book, formula_book)
    ir = IR(
        source_file=path.name,
        file_hash=file_hash,
        file_type=path.suffix.lstrip(".").lower(),
        sheets=value_book.sheetnames,
        tables=[],
        notes=[f"FORMULA {formula_cell_to_coord(key)}= {formula_book[key[0]].cell(row=key[1], column=key[2]).value}"
               for key in sorted(solved)],
    )
    for sheet_name in value_book.sheetnames:
        ws = value_book[sheet_name]
        for row in ws.iter_rows():
            cells = [
                CellValue(
                    row=cell.row,
                    col=cell.column,
                    value=_cell_value(cell.value, solved, sheet_name, cell.row, cell.column),
                )
                for cell in row
            ]
            table_row = TableRow(sheet=sheet_name, row_number=row[0].row if row else 0, cells=cells)
            if not table_row.is_empty():
                ir.tables.append(table_row)
    value_book.close()
    formula_book.close()
    if unresolved:
        ir.notes.append(
            "FORMULA 未求值：" + "；".join(f"{sheet}!{coord} {reason}" for sheet, coord, reason in unresolved)
        )
    return ir


def _cell_value(cached, solved: dict[tuple[str, int, int], float], sheet: str, row: int, col: int):
    """单元格取值：缓存值优先，缓存为空但公式可求值时用求值结果（见 app.formula_eval）。

    浮点值统一按 Excel 显示精度抹掉二进制噪声（0.7000000000000001 → 0.7）：IR 是给
    模型看的"原文"，尾数噪声既让金额出处比对失真，也会把噪声带进下游展示。
    """
    if cached is None:
        value = solved.get((sheet, row, col))
        return display_number(value) if isinstance(value, float) else value
    if isinstance(cached, float):
        return display_number(cached)
    return cached



def _to_ir(path: Path, file_hash: str) -> IR:
    """按扩展名分派到对应 IR 构造器。"""
    from . import ingest_formats

    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return excel_to_ir(path, file_hash)
    if suffix == ".docx":
        return ingest_formats.docx_to_ir(path, file_hash)
    if suffix == ".pdf":
        return ingest_formats.pdf_to_ir(path, file_hash)
    if suffix in (".png", ".jpg", ".jpeg"):
        return ingest_formats.image_to_ir(path, file_hash)
    raise ValueError(f"不支持的文件格式：{path.suffix}（支持 xlsx/xlsm/docx/pdf/png/jpg/jpeg）")


def ingest_file(path: Path, force: bool = False) -> dict:
    """查重 → 接入 → 归档 + IR 快照。返回处理结果描述（命中历史则复用）。"""
    path = path.resolve()
    if not path.exists():
        raise FileNotFoundError(path)

    init_db()
    file_hash = sha256_of(path)
    existing = find_source_file(file_hash)
    if existing and not force:
        return {
            "status": "reused",
            "sha256": file_hash,
            "message": "文件已解析过，复用历史结果",
            "ir_path": existing["ir_path"],
            "archived_path": existing["archived_path"],
        }

    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件格式：{path.suffix}（支持 xlsx/xlsm/docx/pdf/png/jpg/jpeg）")

    ir = _to_ir(path, file_hash)

    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    IR_DIR.mkdir(parents=True, exist_ok=True)
    archived_path = ARCHIVE_DIR / f"{file_hash[:8]}_{path.name}"
    if not archived_path.exists():
        shutil.copy2(path, archived_path)
    ir_path = IR_DIR / f"{file_hash[:8]}.json"
    ir_path.write_text(json.dumps(ir.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")

    conn = get_connection()
    try:
        with conn:
            if existing:
                conn.execute(
                    "UPDATE source_file SET archived_path=?, ir_path=?, original_name=? WHERE sha256=?",
                    (str(archived_path), str(ir_path), path.name, file_hash),
                )
                source_id = existing["id"]
            else:
                cur = conn.execute(
                    """INSERT INTO source_file (sha256, original_name, file_type, archived_path, ir_path)
                       VALUES (?, ?, ?, ?, ?)""",
                    (file_hash, path.name, ir.file_type, str(archived_path), str(ir_path)),
                )
                source_id = cur.lastrowid
    finally:
        conn.close()

    return {
        "status": "ingested",
        "sha256": file_hash,
        "source_id": source_id,
        "sheets": ir.sheets,
        "rows": len(ir.tables),
        "archived_path": str(archived_path),
        "ir_path": str(ir_path),
    }
