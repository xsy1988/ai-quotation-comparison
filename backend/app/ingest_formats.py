"""全格式 IR 构造：docx / pdf（文字层 + 扫描版渲染 OCR）/ 图片（vision OCR）→ 与 excel_to_ir 同构的 IR。

OCR 走 LLM 网关 vision 模型（qwen3.5-ocr），失败上抛不降级（用户政策：不降级，中止并反馈）。
扫描版 PDF 由 pypdfium2 按页渲染为 PNG（200 DPI）后逐页 OCR，sheet 记为 page_N。
qwen3.5-ocr 是专用 OCR 模型，输出契约为原生 {"lines": [{"rotate_rect": [...], "text": ...}]}
（提示词要求的管道格式拿不到、且要求管道格式时表格单元格会被漏检），因此提示词直接对齐原生
契约，脚本侧聚类成 IR 行（sheet|行号|列号:值|...）；模型若直接返回管道字符串也兼容。
聚类前先把窄列竖排表头碎片（同 x 区间、上下紧贴的非数值文本，如"镭雕/破氧/白"）合并为完整
列名条目；行聚类按锚定行高（≤中位行高）截断，高表头单元格不会把下方数据行链进同一行。
IR 行格式与版面理解序列化一致，layout_understand 的 prompt 无需任何改动。
"""

import base64
import io
import json
import re
from pathlib import Path

import pdfplumber
import pypdfium2 as pdfium
from docx import Document

from .ir import CellValue, IR, TableRow, TextBlock
from .prompts import get_prompt

OCR_MODEL = "qwen3.5-ocr"

# 兼容网关忽略 response_format 时返回的裸 IR 行文本
_IR_LINE_RE = re.compile(
    r"^(?P<sheet>[A-Za-z0-9_\u4e00-\u9fff]+)\|(?P<row>\d+)\|(?P<cells>.*)$"
)
_CELL_RE = re.compile(r"^(?P<col>\d+):(?P<value>.*)$")


_CJK_RE = re.compile(r"[一-鿿]")


def _join_fragments(parts) -> str:
    """碎片拼接为完整文本：CJK 边界直接相连（"破氧"+"白"→"破氧白"），拉丁边界补空格。"""
    out = ""
    for part in parts:
        part = str(part).strip()
        if not part:
            continue
        if out and not (_CJK_RE.search(out[-1]) or _CJK_RE.search(part[0])):
            out += " "
        out += part
    return out


def _join_cell_lines(text) -> str:
    """单元格内换行合并为完整文本（pdfplumber 多行表头 "镭雕\n破氧\n白" → "镭雕破氧白"）。

    换行符若带进 IR 会在 sheet|行号|列号:值 序列化中把一行拆成多行孤儿文本，必须在此消除。
    """
    return _join_fragments(str(text).splitlines())


def docx_to_ir(path: Path, file_hash: str) -> IR:
    """python-docx：段落 → TextBlock(sheet="doc", 顺序号)；表格 → TableRow(sheet="table_N")。"""
    doc = Document(path)
    ir = IR(source_file=path.name, file_hash=file_hash, file_type="docx", sheets=["doc"], tables=[])
    seq = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            seq += 1
            ir.blocks.append(TextBlock(text=text, sheet="doc", row=seq, col=1))
    for t_idx, table in enumerate(doc.tables, start=1):
        sheet = f"table_{t_idx}"
        if sheet not in ir.sheets:
            ir.sheets.append(sheet)
        for r_idx, row in enumerate(table.rows, start=1):
            cells = [
                CellValue(row=r_idx, col=c_idx, value=str(cell.text).strip())
                for c_idx, cell in enumerate(row.cells, start=1)
                if str(cell.text).strip()
            ]
            table_row = TableRow(sheet=sheet, row_number=r_idx, cells=cells)
            if not table_row.is_empty():
                ir.tables.append(table_row)
    return ir


def pdf_to_ir(path: Path, file_hash: str, chat_fn=None) -> IR:
    """pdfplumber 逐页：extract_tables → TableRow(sheet=page_N)；页面纯文本 → TextBlock(sheet=page_N)。

    某页文字层与表格都抽不到内容（扫描页）→ pypdfium2 渲染该页为 PNG（200 DPI）→ vision OCR，
    结果归入同一 page_N sheet。LLM 错误上抛不降级；渲染失败抛 ParseError。
    """
    from .ingest import ParseError

    ir = IR(source_file=path.name, file_hash=file_hash, file_type="pdf", sheets=[], tables=[])
    pdf = pdfium.PdfDocument(str(path))
    try:
        page_count = len(pdf)
        with pdfplumber.open(path) as pl:
            for page_no in range(1, page_count + 1):
                sheet = f"page_{page_no}"
                page = pl.pages[page_no - 1]
                page_has_content = False

                for table in page.extract_tables() or []:
                    rows_added = 0
                    for r_idx, row in enumerate(table, start=1):
                        cells = [
                            CellValue(row=r_idx, col=c_idx, value=_join_cell_lines(v))
                            for c_idx, v in enumerate(row, start=1)
                            if v is not None and str(v).strip()
                        ]
                        table_row = TableRow(sheet=sheet, row_number=r_idx, cells=cells)
                        if not table_row.is_empty():
                            ir.tables.append(table_row)
                            rows_added += 1
                    if rows_added:
                        page_has_content = True

                text = (page.extract_text() or "").strip()
                if text:
                    ir.blocks.append(TextBlock(text=text, sheet=sheet, row=1, col=1))
                    page_has_content = True

                if not page_has_content:
                    b64 = _render_pdf_page_b64(pdf, page_no - 1, ParseError)
                    tables, blocks = _ocr_image_b64(b64, "png", sheet, chat_fn)
                    ir.tables.extend(tables)
                    ir.blocks.extend(blocks)
                    page_has_content = bool(tables or blocks)

                if page_has_content:
                    ir.sheets.append(sheet)
    finally:
        pdf.close()
    return ir


def _render_pdf_page_b64(pdf: "pdfium.PdfDocument", index: int, error_cls) -> str:
    """pypdfium2 渲染单页为 200 DPI PNG → base64。渲染失败抛中文 ParseError。"""
    page = pdf[index]
    try:
        try:
            bitmap = page.render(scale=200 / 72)
            pil_image = bitmap.to_pil()
        except Exception as exc:
            raise error_cls(f"扫描版 PDF 第 {index + 1} 页渲染失败：{exc}") from exc
    finally:
        page.close()
    buf = io.BytesIO()
    pil_image.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("ascii")


# 数值类文本不参与竖排碎片合并：避免表头碎片链吞下数据行的金额（如"白"+"0.40"）
_NUMERIC_GUARD_RE = re.compile(r"[0-9.%]")


def _merge_stacked_fragments(entries: list) -> list:
    """窄列竖排换行单元格的碎片合并：x 区间重叠、上下紧贴、均为非数值文本的相邻条目
    按从上到下拼接为完整列名单元（"镭雕"/"破氧"/"白" → "镭雕破氧白"），框取外接矩形。

    entries 为 (x, y, h, w, text)。数值守卫（含数字/% 不合并）保证数据行金额不会被并进来。
    """
    merged: list = []
    for e in sorted(entries, key=lambda it: (it[0], it[1])):
        x, y, h, w, text = e
        if merged:
            px, py, ph, pw, ptext = merged[-1]
            v_gap = y - (py + ph)
            x_overlap = min(px + pw, x + w) - max(px, x)
            if (
                not _NUMERIC_GUARD_RE.search(ptext + text)
                and -0.3 * min(ph, h) <= v_gap <= max(2.0, 0.3 * min(ph, h))
                and x_overlap >= 0.6 * min(pw, w)
            ):
                merged[-1] = (
                    min(px, x), py, (y + h) - py,
                    max(px + pw, x + w) - min(px, x),
                    _join_fragments([ptext, text]),
                )
                continue
        merged.append(e)
    return merged


def _cluster_native_lines(items: list) -> str:
    """qwen3.5-ocr 原生输出（rotate_rect/text 条目）聚类成视觉行，展开为 IR 行文本。

    OCR 专用模型的输出契约不受提示词约束：提示词要求的管道格式拿不到，模型固定返回
    {"lines": [{"rotate_rect": [x, y, h, w, angle], "text": ...}]}（裸 JSON 数组也可能），
    且 angle 恒为 90°——按未旋转框理解：x/y 为左上角，第三、四分量是文本高/宽。
    这里脚本侧确定性转换，不追加 LLM 调用。
    先合并竖排表头碎片（_merge_stacked_fragments）；行聚类以行首条目为锚、
    接受窗高封顶为中位行高（y 区间重叠容差 gap），多行高表头单元格不会把下方
    数据行链进同一视觉行，长文本行也不会因框高被并入邻行。
    """
    entries = []
    for it in items:
        if not isinstance(it, dict) or not isinstance(it.get("text"), str):
            continue
        rect = it.get("rotate_rect") or it.get("bbox") or []
        if isinstance(rect, (list, tuple)) and len(rect) >= 4:
            x, y, h, w = float(rect[0]), float(rect[1]), float(rect[2]), float(rect[3])
        else:
            x, y, h, w = 0.0, float(len(entries)), 10.0, 10.0
        text = _join_cell_lines(it["text"])
        if not text:
            continue
        entries.append((x, y, h, w, text))
    if not entries:
        raise ValueError("OCR 原生 lines 中没有可用文本")
    entries = _merge_stacked_fragments(entries)
    heights = sorted(e[2] for e in entries)
    median_h = heights[len(heights) // 2]
    gap = max(4.0, 0.3 * median_h)
    entries.sort(key=lambda e: e[1])
    rows: list = []  # [y_start, accept_until, [entries]]
    for e in entries:
        if rows and e[1] <= rows[-1][1]:
            rows[-1][2].append(e)
        else:
            # 锚定接受窗：行首条目顶部 + min(自身框高, 中位行高) + gap。
            # 不用运行最大值 y_end，避免竖排/多行高表头单元格逐段链式吞并下方数据行。
            rows.append([e[1], e[1] + min(e[2], median_h) + gap, [e]])

    # 全局列聚类：所有条目的 x 左缘按间距聚成统一列号。
    # 表头两行堆叠的单元格（如"脱墨/氧化"）由此与表体同列对齐，避免版面 LLM 对错列。
    col_gap = max(10.0, 0.6 * median_h)
    col_reps: list[float] = []
    for x in sorted(e[0] for e in entries):
        if col_reps and x - col_reps[-1] <= col_gap:
            col_reps[-1] = (col_reps[-1] + x) / 2
        else:
            col_reps.append(x)

    def col_no(x: float) -> int:
        return min(range(len(col_reps)), key=lambda i: abs(col_reps[i] - x)) + 1

    out_lines = []
    for row_no, (_, _, cells) in enumerate(rows, start=1):
        cells.sort(key=lambda e: (e[1], e[0]))  # 阅读顺序：先上后下、同行从左到右
        by_col: dict[int, list] = {}
        for e in cells:
            by_col.setdefault(col_no(e[0]), []).append(e[4])
        # 列量化后仍同列的碎片按阅读顺序拼回一个单元格（同一逻辑列，不拆成伪列）
        parts = [f"{c}:{_join_fragments(by_col[c])}" for c in sorted(by_col)]
        out_lines.append(f"ocr|{row_no}|" + "|".join(parts))
    return "\n".join(out_lines)


def _decode_ocr_lines(parsed: dict | list) -> str:
    """从 chat_json 返回的 JSON 取 lines；兼容各类退化形态。

    优先 parsed["lines"] 字符串（提示词要求的管道格式）；其次是 OCR 模型原生
    lines 数组（rotate_rect/text，或顶层就是条目数组），聚类转为管道格式；
    最后把 dict 里第一个字符串值当 lines（嵌套 JSON 字符串再解一次）。
    """
    if isinstance(parsed, list):
        return _cluster_native_lines(parsed)
    lines = parsed.get("lines")
    if isinstance(lines, str):
        return lines
    if isinstance(lines, list):
        return _cluster_native_lines(lines)
    for value in parsed.values():
        if isinstance(value, str):
            try:
                inner = json.loads(value)
            except (json.JSONDecodeError, TypeError):
                return value
            if isinstance(inner, dict) and isinstance(inner.get("lines"), str):
                return inner["lines"]
            return value
    raise ValueError("OCR 输出缺少 lines 字段，无法解析 IR 行")


def _parse_ir_lines(lines_text: str) -> tuple[list[TableRow], list]:
    """OCR 文本逐行解析为 TableRow；解析不了的行退化为 TextBlock（顺序保留）。"""
    tables: list[TableRow] = []
    blocks: list[TextBlock] = []
    seq = 0
    for raw in lines_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        seq += 1
        m = _IR_LINE_RE.match(line)
        cells: list[CellValue] = []
        if m:
            for part in m.group("cells").split("|"):
                cm = _CELL_RE.match(part)
                if cm:
                    cells.append(CellValue(row=int(m.group("row")), col=int(cm.group("col")), value=cm.group("value")))
        if cells:
            tables.append(
                TableRow(sheet=m.group("sheet"), row_number=int(m.group("row")), cells=cells)
            )
        else:
            blocks.append(TextBlock(text=line, sheet="ocr", row=seq, col=1))
    return tables, blocks


def _ocr_image_b64(
    b64: str, mime: str, sheet: str, chat_fn=None
) -> tuple[list[TableRow], list[TextBlock]]:
    """base64 图片 → vision OCR（chat_json, model=qwen3.5-ocr）→ IR 行，sheet 归一到调用方指定值。

    LLMError/LLMUnavailable 上抛不捕获（不降级）。OCR 行格式与版面理解 IR 序列化一致。
    """
    from .llm import client as llm_client

    fn = chat_fn or llm_client.chat_json
    messages = [
        {
            "role": "user",
            "content": [
                {"type": "text", "text": get_prompt("ocr_transcribe")},
                {"type": "image_url", "image_url": {"url": f"data:image/{mime};base64,{b64}"}},
            ],
        }
    ]
    parsed, _usage = fn(messages, model=OCR_MODEL)
    tables, blocks = _parse_ir_lines(_decode_ocr_lines(parsed))
    # 提示词固定输出 sheet "ocr"，这里替换为调用方指定的 sheet（page_N / ocr）
    for t in tables:
        t.sheet = sheet
    for b in blocks:
        b.sheet = sheet
    return tables, blocks


def image_to_ir(path: Path, file_hash: str, chat_fn=None) -> IR:
    """图片 → base64 → vision OCR（sheet="ocr"）→ IR。"""
    suffix = path.suffix.lstrip(".").lower()
    mime = "jpeg" if suffix in ("jpg", "jpeg") else suffix
    b64 = base64.b64encode(path.read_bytes()).decode("ascii")
    tables, blocks = _ocr_image_b64(b64, mime, "ocr", chat_fn)
    return IR(
        source_file=path.name,
        file_hash=file_hash,
        file_type="image",
        sheets=["ocr"],
        blocks=blocks,
        tables=tables,
    )
